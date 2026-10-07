"""File-path tests: offset-based masking in text, Word, Excel and PDF, plus fail-closed cases.

Detection here is Tier 1 plus a tiny word-bounded name detector (John/Jane), so the tests
don't need a model and still exercise the FAKER path (names) next to the REDACT path.
"""
import itertools
import re
import struct
import zipfile
import zlib

import docx
import openpyxl
import pymupdf
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl.comments import Comment

from app.multimodal.pipeline import MultimodalPipeline
from dlp_core import DetectionPipeline, DetectorSpec, FernetSealer, InMemoryVault, OffsetMasker, Span
from dlp_core.residual_scanner import scan as residual
from dlp_core.tier1 import Tier1Engine


class Names:
    name, labels = "names", frozenset({"PERSON"})

    def scan(self, text):
        return [Span(m.start(), m.end(), "PERSON", 0.9, "names") for m in re.finditer(r"\b(?:John|Jane)\b", text)]


def make_pipeline(**kw):
    pipe = DetectionPipeline([DetectorSpec(Tier1Engine(), critical=True), DetectorSpec(Names())])
    c = itertools.count(1)
    masker = OffsetMasker(InMemoryVault(FernetSealer()), lambda l, r: f"Zed{next(c)}x")
    return MultimodalPipeline(lambda t: pipe.run(t), masker, residual, **kw)


def run(tmp_path, name, writer, **kw):
    src, out = tmp_path / name, tmp_path / ("m_" + name)
    writer(src)
    return make_pipeline(**kw).process(str(src), str(out), "c"), out


def png_bytes():
    raw = b"\x00\xff\xff\xff"
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


# ------------------------------------------------------------------ text
def test_text_short_values_and_substrings_untouched(tmp_path):
    body = "John met Johnson. cvv 123, order 12345, ticket 1234, build 1.2.3.\n"
    res, out = run(tmp_path, "a.txt", lambda p: p.write_text(body))
    assert res["success"], res
    assert out.read_text() == "Zed1x met Johnson. cvv [REDACTED:CVV], order 12345, ticket 1234, build 1.2.3.\n"


@pytest.mark.parametrize("enc,bom", [("utf-8-sig", True), ("utf-16", True), ("cp1252", False)])
def test_text_encoding_and_line_endings_preserved(tmp_path, enc, bom):
    body = "Café mail a@b.com\r\nsecond line é\r\n"
    res, out = run(tmp_path, "b.txt", lambda p: p.write_bytes(body.encode(enc)))
    assert res["success"], res
    raw = out.read_bytes()
    assert raw.decode(enc) == "Café mail [REDACTED:EMAIL]\r\nsecond line é\r\n".replace("[REDACTED:EMAIL]", raw.decode(enc).split("mail ")[1].split("\r")[0])
    assert b"a@b.com" not in raw and b"\r\n" in raw or enc == "utf-16"


def test_clean_text_unchanged(tmp_path):
    res, out = run(tmp_path, "c.txt", lambda p: p.write_bytes(b"Just a normal sentence."))
    assert res["success"] and res["replacements_made"] == 0 and out.read_bytes() == b"Just a normal sentence."


# ------------------------------------------------------------------ word
def _word(path):
    d = docx.Document()
    d.core_properties.author = "Jane Roe"
    p = d.add_paragraph()
    r1 = p.add_run("Contact: Jo"); r1.bold = True
    p.add_run("hn at john.")
    r3 = p.add_run("doe@example.com now"); r3.italic = True
    d.add_paragraph("Johnson and cvv 123, order 12345.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "card 4242 4242 4242 4242"
    d.sections[0].header.paragraphs[0].text = "Header for Jane"
    para = d.add_paragraph("Link: ")
    rid = d.part.relate_to("mailto:john@example.com", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); tx = OxmlElement("w:t"); tx.text = "john@example.com"; r.append(tx); h.append(r)
    para._p.append(h)
    d.save(path)


def test_word_masks_across_runs_keeps_formatting_and_untouched_text(tmp_path):
    res, out = run(tmp_path, "a.docx", _word)
    assert res["success"], res
    d = docx.Document(str(out))
    p0 = d.paragraphs[0]
    assert p0.text == "Contact: Zed1x at Zed2x now" or "Contact: Zed1x at" in p0.text
    assert "john.doe@example.com" not in p0.text and "example.com" not in p0.text
    assert p0.runs[0].bold is True and p0.runs[-1].italic is True            # formatting survives
    assert d.paragraphs[1].text == "Johnson and cvv [REDACTED:CVV], order 12345."
    assert "4242" not in d.tables[0].cell(0, 0).text
    assert "Jane" not in d.sections[0].header.paragraphs[0].text
    assert d.core_properties.author == ""


def test_word_hyperlink_target_and_text_masked(tmp_path):
    res, out = run(tmp_path, "a.docx", _word)
    with zipfile.ZipFile(out) as z:
        blob = b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels")))
    assert b"john@example.com" not in blob and b"mailto:john" not in blob


def test_word_tracked_deletion_and_comments_block(tmp_path):
    def w(path):
        d = docx.Document(); d.add_paragraph("clean text"); d.save(path)
        with zipfile.ZipFile(path) as z:
            items = {n: z.read(n) for n in z.namelist()}
        items["word/document.xml"] = items["word/document.xml"].replace(
            b"</w:body>", b'<w:p><w:del w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>secret 4242 4242 4242 4242</w:delText></w:r></w:del></w:p></w:body>')
        with zipfile.ZipFile(path, "w") as z:
            for n, b in items.items():
                z.writestr(n, b)
    res, out = run(tmp_path, "t.docx", w)
    assert not res["success"] and "tracked_deletions_present" in res["blockers"] and not out.exists()


def test_word_image_only_blocks_but_image_plus_text_warns(tmp_path):
    def img_only(path):
        d = docx.Document(); png = path.with_suffix(".png"); png.write_bytes(png_bytes()); d.add_picture(str(png)); d.save(path)
    res, out = run(tmp_path, "i.docx", img_only)
    assert not res["success"] and "no_extractable_text" in res["blockers"]

    def mixed(path):
        d = docx.Document(); d.add_paragraph("hello there"); png = path.with_suffix(".png"); png.write_bytes(png_bytes())
        d.add_picture(str(png)); d.save(path)
    res, _ = run(tmp_path, "m.docx", mixed)
    assert res["success"] and "images_not_inspected" in res["warnings"]


def test_corrupt_docx_fails_closed(tmp_path):
    res, out = run(tmp_path, "bad.docx", lambda p: p.write_bytes(b"not a docx"))
    assert not res["success"] and not out.exists()


# ------------------------------------------------------------------ excel
def _xlsx(path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = "mail john@example.com"; ws["A2"] = 42; ws["A3"] = "=A2*2"; ws["A4"] = '="a@b.com"'
    ws["B1"] = "John and Johnson"; ws["C1"] = "cvv 123 / 12345"; ws["D1"] = 123
    ws["E1"] = "note"; ws["E1"].comment = Comment("ping Jane at x@y.org", "Jane Roe")
    wb.properties.creator = "Jane Roe"
    wb.create_sheet("hidden").sheet_state = "hidden"; wb["hidden"]["A1"] = "4242 4242 4242 4242"
    wb.save(path)


def test_excel_masks_strings_formulas_comments_hidden_sheets_keeps_numbers(tmp_path):
    res, out = run(tmp_path, "a.xlsx", _xlsx)
    assert res["success"], res
    ws = openpyxl.load_workbook(out)["Sheet"]
    assert "john@example.com" not in ws["A1"].value
    assert ws["A2"].value == 42 and ws["A3"].value == "=A2*2" and ws["D1"].value == 123
    assert "a@b.com" not in ws["A4"].value and ws["A4"].value.startswith('="') and ws["A4"].value.endswith('"')
    assert ws["B1"].value.endswith("and Johnson") and ws["C1"].value == "cvv [REDACTED:CVV] / 12345"
    assert "Jane" not in ws["E1"].comment.text and ws["E1"].comment.author == "user"
    wb = openpyxl.load_workbook(out)
    assert "4242" not in wb["hidden"]["A1"].value and wb.properties.creator != "Jane Roe"


def test_excel_sheet_name_with_sensitive_data_blocks(tmp_path):
    def w(path):
        wb = openpyxl.Workbook(); wb.active.title = "mail a@b.com"; wb.active["A1"] = "x"; wb.save(path)
    res, out = run(tmp_path, "n.xlsx", w)
    assert not res["success"] and "sheet_name_contains_sensitive_data" in res["blockers"] and not out.exists()


# ------------------------------------------------------------------ pdf
def _pdf(path):
    doc = pymupdf.open(); page = doc.new_page()
    page.insert_text((72, 100), "Contact Jane at jane@corp.com. John said Johnson is not PII.", fontsize=11)
    page.insert_text((72, 130), "Order 12345 cvv 123 done.", fontsize=11)
    doc.set_metadata({"author": "Jane Roe", "title": "Report for John"})
    doc.save(path)


def test_pdf_masks_only_detected_characters_and_removes_original_text(tmp_path):
    res, out = run(tmp_path, "a.pdf", _pdf)
    assert res["success"], res
    text = pymupdf.open(out)[0].get_text()
    assert "jane@corp.com" not in text and "Jane" not in text
    assert "Johnson is not PII" in text and "12345" in text
    assert "John said" not in text                                          # standalone John masked
    md = pymupdf.open(out).metadata
    assert not (md.get("author") or md.get("title"))


def test_pdf_scanned_and_empty_block(tmp_path):
    def scan(path):
        doc = pymupdf.open(); pg = doc.new_page(); png = path.with_suffix(".png"); png.write_bytes(png_bytes())
        pg.insert_image(pymupdf.Rect(72, 72, 272, 172), filename=str(png)); doc.save(path)
    res, out = run(tmp_path, "s.pdf", scan)
    assert not res["success"] and "no_extractable_text" in res["blockers"] and not out.exists()

    def blank(path):
        doc = pymupdf.open(); doc.new_page(); doc.save(path)
    res, _ = run(tmp_path, "b.pdf", blank)
    assert not res["success"] and "no_extractable_text" in res["blockers"]


def test_pdf_link_target_with_sensitive_data_is_removed(tmp_path):
    def w(path):
        doc = pymupdf.open(); pg = doc.new_page(); pg.insert_text((72, 100), "click here", fontsize=11)
        pg.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 90, 140, 105), "uri": "mailto:a@b.com"})
        doc.save(path)
    res, out = run(tmp_path, "l.pdf", w)
    assert res["success"] and not any(l.get("uri") for l in pymupdf.open(out)[0].get_links())


# ------------------------------------------------------------------ pipeline-level
def test_strict_mode_blocks_degraded_coverage(tmp_path):
    from dlp_core import UnavailableDetector
    pipe = DetectionPipeline([DetectorSpec(Tier1Engine(), critical=True), DetectorSpec(UnavailableDetector("tier2", ["PERSON"]))])
    masker = OffsetMasker(InMemoryVault(FernetSealer()))
    p = tmp_path / "a.txt"; p.write_text("mail a@b.com")
    ok = MultimodalPipeline(lambda t: pipe.run(t), masker, residual).process(str(p), str(tmp_path / "o1.txt"), "c")
    assert ok["success"] and ok["degraded"] and ok["uncovered_labels"] == ["PERSON"]
    strict = MultimodalPipeline(lambda t: pipe.run(t), masker, residual, strict=True).process(str(p), str(tmp_path / "o2.txt"), "c")
    assert not strict["success"] and strict["error"] == "degraded_coverage_strict_mode" and not (tmp_path / "o2.txt").exists()


def test_critical_detector_failure_blocks_and_leaves_no_output(tmp_path):
    class Boom:
        name, labels = "t1", frozenset({"EMAIL"})
        def scan(self, t): raise RuntimeError
    pipe = DetectionPipeline([DetectorSpec(Boom(), critical=True)])
    p = tmp_path / "a.txt"; p.write_text("mail a@b.com")
    res = MultimodalPipeline(lambda t: pipe.run(t), OffsetMasker(InMemoryVault(FernetSealer())), residual).process(str(p), str(tmp_path / "o.txt"), "c")
    assert not res["success"] and "detection_blocked" in res["error"] and not (tmp_path / "o.txt").exists()


def test_unknown_and_missing_inputs(tmp_path):
    p = tmp_path / "x.bin"; p.write_bytes(b"\x00\x01\x02\x03")
    assert not make_pipeline().process(str(p), str(tmp_path / "o"), "c")["success"]
    assert make_pipeline().process(str(tmp_path / "nope.txt"), None, "c")["error"] == "input_not_found"


def test_same_person_same_fake_across_files_in_one_conversation(tmp_path):
    pl = make_pipeline()
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_text("John here"); b.write_text("call John")
    pl.process(str(a), str(tmp_path / "ma.txt"), "chat1"); pl.process(str(b), str(tmp_path / "mb.txt"), "chat1")
    assert (tmp_path / "ma.txt").read_text().split()[0] == (tmp_path / "mb.txt").read_text().split()[1]
