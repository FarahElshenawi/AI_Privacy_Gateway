"""Table context, numeric Excel cells, content-based file typing, PEM with literal \\n."""
from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.multimodal.file_type_detector import detect_file_type
from app.multimodal.pipeline import MultimodalPipeline
from dlp_core.segments import Segment, SegmentMasker
from dlp_core.tier1.validators import pem_valid


def _run(path, out):
    return MultimodalPipeline().process(str(path), str(out))


# --- tabular context -----------------------------------------------------------

def test_xlsx_numeric_cells_with_header_context_are_masked(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["Name", "SSN", "Card", "Account"])
    ws.append(["Alice", "123-45-6789", 4111111111111111, "987654321"])
    src = tmp_path / "a.xlsx"; wb.save(src)
    r = _run(src, tmp_path / "a_m.xlsx")
    assert r["success"], r
    row = [c.value for c in load_workbook(tmp_path / "a_m.xlsx").active[2]]
    assert "[REDACTED:US_SSN]" in row and "[REDACTED:CREDIT_CARD]" in row
    assert "4111111111111111" not in " ".join(map(str, row)) and "987654321" not in row
    assert [c.value for c in load_workbook(tmp_path / "a_m.xlsx").active[1]][1] == "SSN"   # header untouched


def test_xlsx_non_sensitive_numbers_stay_numbers(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["Item", "Qty", "Price"])
    ws.append(["Widget", 3, 9.5])
    src = tmp_path / "n.xlsx"; wb.save(src)
    assert _run(src, tmp_path / "n_m.xlsx")["success"]
    row = [c.value for c in load_workbook(tmp_path / "n_m.xlsx").active[2]]
    assert row == ["Widget", 3, 9.5]


def test_context_is_never_written_back():
    from app.pipeline.engine import _masker, detect
    sm = SegmentMasker(detect, _masker, max_batch_chars=10_000)
    res = sm.mask_segments([Segment("a", "123-45-6789", "SSN: ")], "ctx-test")
    assert res.masked["a"] == "[REDACTED:US_SSN]"      # the "SSN: " context is not in the output
    assert all(e.start >= 0 and e.end <= len("123-45-6789") for e in res.edits["a"])


def test_csv_cells_get_header_context_and_layout_is_preserved(tmp_path):
    src = tmp_path / "p.csv"
    src.write_text('Name,SSN,Note\r\nAlice,123-45-6789,"hi, 4111 1111 1111 1111"\r\nBob,x,y\r\n')
    r = _run(src, tmp_path / "p_m.csv")
    assert r["success"], r
    out = (tmp_path / "p_m.csv").read_bytes().decode()
    assert out.startswith("Name,SSN,Note\r\n") and out.endswith("Bob,x,y\r\n")
    assert "123-45-6789" not in out and "4111" not in out and '"hi, [REDACTED:CREDIT_CARD]"' in out


def test_csv_quoted_newline_stays_in_one_cell(tmp_path):
    src = tmp_path / "q.csv"
    src.write_text('A,B\n1,"line1\nSSN 123-45-6789"\n')
    r = _run(src, tmp_path / "q_m.csv")
    assert r["success"] and "123-45-6789" not in (tmp_path / "q_m.csv").read_text()


def test_docx_table_cells_use_the_column_header(tmp_path):
    from docx import Document
    d = Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "SSN"
    t.cell(1, 0).text, t.cell(1, 1).text = "Alice", "123-45-6789"
    src = tmp_path / "t.docx"; d.save(src)
    r = _run(src, tmp_path / "t_m.docx")
    assert r["success"], r
    out = Document(tmp_path / "t_m.docx").tables[0]
    assert out.cell(1, 1).text.startswith("[REDACTED") and out.cell(0, 1).text == "SSN"


# --- file type by content ------------------------------------------------------

def _zip(path, names):
    with zipfile.ZipFile(path, "w") as z:
        for n in names:
            z.writestr(n, "<x/>")


def test_zip_content_decides_the_type_not_the_name(tmp_path):
    _zip(tmp_path / "a.docx", ["xl/workbook.xml", "[Content_Types].xml"])
    _zip(tmp_path / "b.xlsx", ["word/document.xml", "[Content_Types].xml"])
    _zip(tmp_path / "c.txt", ["word/document.xml"])
    _zip(tmp_path / "d.docx", ["something/else.xml"])
    assert detect_file_type(str(tmp_path / "a.docx")) == "excel"
    assert detect_file_type(str(tmp_path / "b.xlsx")) == "word"
    assert detect_file_type(str(tmp_path / "c.txt")) == "word"
    assert detect_file_type(str(tmp_path / "d.docx")) == "unknown"


@pytest.mark.parametrize("name,data", [
    ("fake.pdf", b"hello"), ("fake.docx", b"hello"), ("fake.xlsx", b"hello"),
    ("old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 20), ("blob.bin", b"\x00\x01\x02\x03"),
])
def test_names_that_lie_or_unsupported_binaries_are_unknown(tmp_path, name, data):
    (tmp_path / name).write_bytes(data)
    assert detect_file_type(str(tmp_path / name)) == "unknown"


def test_real_files_and_text_variants_are_recognised(tmp_path):
    (tmp_path / "x.pdf").write_bytes(b"%PDF-1.7\n")
    (tmp_path / "noext").write_text("just text, no extension")
    (tmp_path / "u16.txt").write_bytes("héllo".encode("utf-16"))
    (tmp_path / "l1.txt").write_bytes("caf\xe9 au lait".encode("latin-1"))
    (tmp_path / "i.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 8)
    assert [detect_file_type(str(tmp_path / n)) for n in ("x.pdf", "noext", "u16.txt", "l1.txt", "i.png")] == \
        ["pdf", "text", "text", "text", "image"]


def test_wrongly_named_docx_is_processed_as_what_it_is(tmp_path):
    wb = Workbook(); ws = wb.active
    ws.append(["Name", "SSN"]); ws.append(["Alice", "123-45-6789"])
    buf = io.BytesIO(); wb.save(buf)
    (tmp_path / "report.docx").write_bytes(buf.getvalue())
    r = _run(tmp_path / "report.docx", tmp_path / "out.xlsx")
    assert r["success"] and r["input_type"] == "excel"


# --- PEM with literal \n ---------------------------------------------------------

_B64 = "MIIEvQIBADANBgkqhkiG9w0BAQEFAASCBKcwggSjAgEAAoIBAQC7" * 3


def test_pem_with_escaped_newlines_validates():
    assert pem_valid(f"-----BEGIN PRIVATE KEY-----\\n{_B64}\\n-----END PRIVATE KEY-----\\n")
    assert pem_valid(f"-----BEGIN PRIVATE KEY-----\n{_B64}\n-----END PRIVATE KEY-----")
    assert not pem_valid("-----BEGIN PRIVATE KEY-----\\nnot base64 !!!\\n-----END PRIVATE KEY-----")


def test_gcp_service_account_json_is_masked(tmp_path):
    key = f"-----BEGIN PRIVATE KEY-----\\n{_B64}\\n-----END PRIVATE KEY-----\\n"
    src = tmp_path / "sa.json"
    src.write_text('{"type":"service_account","private_key_id":"abc","private_key":"%s",'
                   '"client_email":"svc@proj.iam.gserviceaccount.com"}' % key)
    r = _run(src, tmp_path / "sa_m.json")
    assert r["success"], r
    out = (tmp_path / "sa_m.json").read_text()
    assert "MIIEvQ" not in out and "BEGIN PRIVATE KEY" not in out and '"type":"service_account"' in out
