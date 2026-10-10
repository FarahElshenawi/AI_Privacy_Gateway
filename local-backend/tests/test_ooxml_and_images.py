"""OOXML metadata scrubbing, uninspectable parts, alt text, and the image policy."""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from app.multimodal import ooxml
from app.multimodal.pipeline import MultimodalPipeline
from app.pipeline import engine

def _make_png() -> bytes:
    import struct
    import zlib

    def chunk(t: bytes, d: bytes) -> bytes:
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = b"\x00\xff\x00\x00\xff"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


PNG = _make_png()


def _run(src, out):
    return MultimodalPipeline().process(str(src), str(out))


def _rewrite(path, edits: dict[str, bytes | None], extra: dict[str, bytes] | None = None):
    """Rewrite a zip: edits maps part -> new bytes (None deletes); extra adds parts."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            if i.filename in edits:
                if edits[i.filename] is not None:
                    zout.writestr(i.filename, edits[i.filename])
            elif i.filename not in (extra or {}):
                zout.writestr(i, zin.read(i.filename))
        for n, d in (extra or {}).items():
            zout.writestr(n, d)
    path.write_bytes(buf.getvalue())


def _docx(tmp_path, name="d.docx"):
    d = Document()
    d.add_paragraph("Quarterly notes for the team.")
    p = tmp_path / name
    d.save(p)
    return p


def test_metadata_is_scrubbed_from_docx(tmp_path):
    src = _docx(tmp_path)
    _rewrite(src, {}, {
        "docProps/custom.xml": b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties">'
                               b'<property name="Client"><lpwstr>Acme Secret Client</lpwstr></property></Properties>',
        "docProps/app.xml": b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                            b'<Company>Initech Bank</Company><Manager>Jane Roe</Manager></Properties>',
        "docProps/thumbnail.png": PNG,
    })
    r = _run(src, tmp_path / "o.docx")
    assert r["success"], r
    with zipfile.ZipFile(tmp_path / "o.docx") as z:
        blob = b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels")))
        assert b"Acme Secret Client" not in blob and b"Initech Bank" not in blob and b"Jane Roe" not in blob
        assert not any("thumbnail" in n for n in z.namelist())
        assert b"thumbnail" not in z.read("_rels/.rels") and b"thumbnail" not in z.read("[Content_Types].xml").lower() \
            .replace(b'extension="png"', b"")
    Document(tmp_path / "o.docx")                       # still opens
    assert "metadata_scrubbed" in r["warnings"]


def test_tracked_change_author_names_are_removed(tmp_path):
    src = _docx(tmp_path)
    with zipfile.ZipFile(src) as z:
        doc = z.read("word/document.xml").decode()
    ins = ('<w:ins w:id="9" w:author="Jane Doe" w:date="2026-01-01T00:00:00Z" w:initials="JD"><w:r><w:t> added</w:t></w:r></w:ins>')
    _rewrite(src, {"word/document.xml": doc.replace("</w:p>", ins + "</w:p>", 1).encode()})
    r = _run(src, tmp_path / "o.docx")
    assert r["success"], r
    with zipfile.ZipFile(tmp_path / "o.docx") as z:
        xml = z.read("word/document.xml").decode()
    assert "Jane Doe" not in xml and 'w:author="user"' in xml and "JD" not in xml


def test_reviewer_list_part_is_dropped_with_its_relationship(tmp_path):
    src = _docx(tmp_path)
    with zipfile.ZipFile(src) as z:
        rels = z.read("word/_rels/document.xml.rels").decode()
        ct = z.read("[Content_Types].xml").decode()
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdP" Type="http://schemas.microsoft.com/office/2011/relationships/people" Target="people.xml"/></Relationships>')
    ct = ct.replace("</Types>", '<Override PartName="/word/people.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml"/></Types>')
    _rewrite(src, {"word/_rels/document.xml.rels": rels.encode(), "[Content_Types].xml": ct.encode()},
             {"word/people.xml": b"<w15:people xmlns:w15='x'><w15:person w15:author='Jane Doe'/></w15:people>"})
    assert _run(src, tmp_path / "o.docx")["success"]
    with zipfile.ZipFile(tmp_path / "o.docx") as z:
        assert "word/people.xml" not in z.namelist()
        assert b"people" not in z.read("word/_rels/document.xml.rels") and b"people" not in z.read("[Content_Types].xml")
    Document(tmp_path / "o.docx")


@pytest.mark.parametrize("part,code", [
    ("word/charts/chart1.xml", "charts_not_inspected"),
    ("word/diagrams/data1.xml", "smartart_not_inspected"),
    ("word/glossary/document.xml", "glossary_document_present"),
])
def test_uninspectable_parts_block_by_default_and_warn_when_configured(tmp_path, monkeypatch, part, code):
    src = _docx(tmp_path)
    _rewrite(src, {}, {part: b"<c:chartSpace xmlns:c='x'><c:v>nothing sensitive</c:v></c:chartSpace>"})
    r = _run(src, tmp_path / "o.docx")
    assert not r["success"] and code in r["blockers"]
    monkeypatch.setenv("DLP_OOXML_PARTS", "warn")
    r2 = _run(src, tmp_path / "o2.docx")
    if code == "glossary_document_present":
        assert not r2["success"]                      # not downgradable: it holds document text
    else:
        assert r2["success"] and code in r2["warnings"]


def test_warn_mode_still_blocks_hard_evidence_inside_the_unscanned_part(tmp_path, monkeypatch):
    monkeypatch.setenv("DLP_OOXML_PARTS", "warn")
    src = _docx(tmp_path)
    with zipfile.ZipFile(src) as z:
        rels = z.read("word/_rels/document.xml.rels").decode()
        ct = z.read("[Content_Types].xml").decode()
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdChart1" Type="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/></Relationships>')
    ct = ct.replace("</Types>", '<Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-'
                    'officedocument.drawingml.chart+xml"/></Types>')
    _rewrite(src, {"word/_rels/document.xml.rels": rels.encode(), "[Content_Types].xml": ct.encode()},
             {"word/charts/chart1.xml": b"<c:v>card 4111 1111 1111 1111</c:v>"})
    r = _run(src, tmp_path / "o.docx")
    assert not r["success"] and r["leaks"]


def test_custom_xml_with_text_blocks(tmp_path):
    src = _docx(tmp_path)
    _rewrite(src, {}, {"customXml/item1.xml": b"<root><client>Acme Secret Client</client></root>"})
    r = _run(src, tmp_path / "o.docx")
    assert not r["success"] and "custom_xml_with_text" in r["blockers"]


def test_alt_text_and_picture_names_are_masked(tmp_path):
    src = _docx(tmp_path)
    with zipfile.ZipFile(src) as z:
        doc = z.read("word/document.xml").decode()
    drawing = ('<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
               '<wp:docPr id="1" name="passport_scan.png" descr="Photo of passport, contact bob@example.com"/></wp:inline></w:drawing></w:r>')
    doc = doc.replace("</w:p>", drawing + "</w:p>", 1)
    if "xmlns:wp=" not in doc.split(">", 2)[1] and "wordprocessingDrawing" not in doc[:2000]:
        pass
    _rewrite(src, {"word/document.xml": doc.encode()})
    r = _run(src, tmp_path / "o.docx")
    assert r["success"], r
    with zipfile.ZipFile(tmp_path / "o.docx") as z:
        xml = z.read("word/document.xml").decode()
    assert "bob@example.com" not in xml and 'name="passport_scan.png"' in xml     # a plain file name is not PII


def test_xlsx_metadata_is_scrubbed(tmp_path):
    wb = Workbook(); wb.active.append(["Name", "Note"]); wb.active.append(["x", "hello"])
    src = tmp_path / "a.xlsx"; wb.save(src)
    _rewrite(src, {}, {"docProps/custom.xml": b"<Properties xmlns='http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'>"
                                              b"<property name='Client'><lpwstr>Acme Secret Client</lpwstr></property></Properties>"})
    r = _run(src, tmp_path / "o.xlsx")
    assert r["success"], r
    with zipfile.ZipFile(tmp_path / "o.xlsx") as z:
        assert b"Acme Secret Client" not in b"".join(z.read(n) for n in z.namelist() if n.endswith(".xml"))


# --- image policy ---------------------------------------------------------------

@pytest.fixture
def image_policy():
    old = engine.IMAGE_POLICY
    yield lambda v: engine.apply_tenant_config([], [], v)
    engine.IMAGE_POLICY = old


def test_standalone_image_default_blocks_and_warn_passes_it_through_unchanged(tmp_path, image_policy):
    src = tmp_path / "p.png"; src.write_bytes(PNG)
    r = _run(src, tmp_path / "o.png")
    assert not r["success"] and "image_file" in r["blockers"]
    image_policy("warn")
    r = _run(src, tmp_path / "o.png")
    assert r["success"] and "image_not_inspected" in r["warnings"] and (tmp_path / "o.png").read_bytes() == PNG
    image_policy("block")
    assert not _run(src, tmp_path / "o2.png")["success"]


def test_images_inside_documents_follow_the_policy(tmp_path, image_policy):
    d = Document(); d.add_paragraph("Some text with a logo.")
    img = tmp_path / "logo.png"; img.write_bytes(PNG)
    d.add_picture(str(img))
    src = tmp_path / "i.docx"; d.save(src)
    r = _run(src, tmp_path / "o.docx")
    assert r["success"] and "images_not_inspected" in r["warnings"]            # default: allowed with a warning
    image_policy("block")
    r = _run(src, tmp_path / "o2.docx")
    assert not r["success"] and "images_present" in r["blockers"]
    image_policy("warn")
    assert _run(src, tmp_path / "o3.docx")["success"]


def test_image_policy_comes_from_the_cloud_tenant_config(image_policy):
    from app import cloud_sync
    cloud_sync._tenant_version = None
    assert cloud_sync.apply_tenant_config({"deny_terms": [], "tenant_domains": [], "image_policy": "block", "version": 1})
    assert engine.IMAGE_POLICY == "block"
    assert not cloud_sync.apply_tenant_config({"image_policy": "ocr"})        # unknown value: ignored
    assert engine.IMAGE_POLICY == "block"
    cloud_sync._tenant_version = None
    engine.apply_tenant_config([], [])


def test_parts_policy_env_defaults_to_block(monkeypatch):
    monkeypatch.delenv("DLP_OOXML_PARTS", raising=False)
    assert ooxml.parts_policy() == "block"
    monkeypatch.setenv("DLP_OOXML_PARTS", "nonsense")
    assert ooxml.parts_policy() == "block"
