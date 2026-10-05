"""Word document parser — extracts paragraphs and runs.

Extracts text from body paragraphs AND table cells, so PII in tables
is detected and masked.
"""
from docx import Document


class WordParser:
    def parse(self, file_path: str) -> dict:
        """Extract text from a Word document, paragraph by paragraph.

        Preserves run-level formatting (bold, italic) so the reconstructor
        can apply the same formatting to the masked text.

        Returns: {'paragraphs': [{'text': str, 'runs': [(text, bold, italic), ...]}, ...]}

        Raises:
            ValueError: if the file is not a valid .docx (corrupt or wrong format)
        """
        try:
            doc = Document(file_path)
        except Exception as e:
            raise ValueError(f"Cannot parse Word document: {e}") from e

        paragraphs = []

        # Body paragraphs
        for para in doc.paragraphs:
            runs = [(run.text, run.bold, run.italic) for run in para.runs]
            paragraphs.append({
                "text": para.text,
                "runs": runs,
            })

        # Table cells (each cell has paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        runs = [(run.text, run.bold, run.italic) for run in para.runs]
                        paragraphs.append({
                            "text": para.text,
                            "runs": runs,
                        })

        # Headers and footers
        for section in doc.sections:
            for para in section.header.paragraphs:
                runs = [(run.text, run.bold, run.italic) for run in para.runs]
                paragraphs.append({"text": para.text, "runs": runs})
            for para in section.footer.paragraphs:
                runs = [(run.text, run.bold, run.italic) for run in para.runs]
                paragraphs.append({"text": para.text, "runs": runs})

        return {"paragraphs": paragraphs}
